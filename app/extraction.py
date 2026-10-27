"""Safe text extraction from uploaded PDF / DOCX files.

Files are only parsed as data (never executed or saved to disk). Scanned, image-only PDFs
are NOT supported (no OCR in this MVP) and produce a clear error.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

ALLOWED_EXTENSIONS = {".pdf", ".docx"}
MIN_TEXT_CHARS = 30


class ExtractionError(Exception):
    """Raised with a recruiter-friendly message when a file cannot be processed."""


@dataclass
class ExtractionResult:
    text: str
    file_type: str
    warnings: list[str] = field(default_factory=list)


def _validate(filename: str, data: bytes, max_bytes: int) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ExtractionError(
            f"Unsupported file type '{ext or 'unknown'}'. Only PDF and DOCX files are accepted."
        )
    if not data:
        raise ExtractionError("The uploaded file is empty (0 bytes).")
    if len(data) > max_bytes:
        raise ExtractionError(
            f"File is too large ({len(data) / 1024 / 1024:.1f} MB). Limit is {max_bytes / 1024 / 1024:.0f} MB."
        )
    # Check magic bytes so a renamed file cannot masquerade as a resume.
    if ext == ".pdf" and not data.lstrip()[:5] == b"%PDF-":
        raise ExtractionError("File has a .pdf extension but is not a valid PDF.")
    if ext == ".docx" and not data[:2] == b"PK":
        raise ExtractionError("File has a .docx extension but is not a valid DOCX.")
    return ext


def extract_pdf(data: bytes) -> str:
    import pymupdf  # imported lazily

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # corrupted PDF
        raise ExtractionError(f"Could not open PDF: {exc}") from exc
    try:
        if doc.needs_pass:
            raise ExtractionError("PDF is password-protected and cannot be read.")
        return "\n".join(page.get_text("text") for page in doc)
    finally:
        doc.close()


def extract_docx(data: bytes) -> str:
    import docx  # python-docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not open DOCX: {exc}") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes, max_bytes: int = 5 * 1024 * 1024) -> ExtractionResult:
    ext = _validate(filename, data, max_bytes)
    text = extract_pdf(data) if ext == ".pdf" else extract_docx(data)
    text = text.replace("\x00", "").strip()
    if len(text) < MIN_TEXT_CHARS:
        hint = (
            " This looks like a scanned/image-only PDF; OCR is not supported in this MVP."
            if ext == ".pdf"
            else ""
        )
        raise ExtractionError("No readable text found in the document." + hint)
    return ExtractionResult(text=text, file_type=ext.lstrip("."))
